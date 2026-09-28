"""Step 9：生成品控指标字典（Excel + Markdown）。

每个指标的"口径 SQL"都会被实际执行一遍，得到 2018 年 1-8 月的基准值写进字典：
字典里的 SQL 不是示意，而是能跑、跑出来就是看板/报告上的那个数。
"""
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common import OUT_DIR, ROOT, query

W = "in_scope = 1 AND purchase_month BETWEEN '2018-01' AND '2018-08'"   # 基准窗口
SC = "（商家评分窗口：2018-03 ~ 2018-08）"

# 字段：编码, 名称, 层级, 类型, 业务域, 父指标, 业务定义, 计算公式, 统计范围/过滤, 下钻维度, 来源表, 口径SQL, 单位, 方向, 目标/预警, 更新频率, 可得性, 易错点
M = [
    # ------------------------------------------------------------------ L0 北极星
    ("QC001", "品质客诉率", "L0 北极星", "结果", "品质", "—",
     "签收订单中，用户因商品本身问题（假货/质量缺陷/货不对板/少件漏发/包装破损）给出 1-3 星并在评价中描述该问题的订单占比。衡量\"用户收到的东西有没有问题\"。",
     "品质客诉签收单数 ÷ 签收订单数",
     "签收订单（status=delivered 且签收时间非空）；按下单月归属；品质客诉只认 1-3 星",
     "月/周、一级类目、品类、商家、省份、订单结构、问题类型",
     "dwd_qc_order",
     f"SELECT SUM(is_delivered * is_quality_complaint) / SUM(is_delivered) FROM dwd_qc_order WHERE {W}",
     "%", "越低越好", "目标 ≤ 4.0%；预警 > 5.5%", "日更（T+1）", "已上线",
     "① 分子分母必须是同一批订单（都按下单月），不能用\"本月评价数 ÷ 本月签收数\"；② 一单只计一次，不按问题类型重复计数；③ 4-5 星里提到\"没坏/无瑕疵\"的不算"),
    # ------------------------------------------------------------------ L1 品质拆解
    ("QC101", "少件/漏发率", "L1 结果-品质拆解", "结果", "品质", "QC001",
     "签收后用户反馈少发、漏发、缺配件的订单占比。",
     "少件漏发客诉签收单 ÷ 签收订单数", "同 QC001", "同 QC001 + 单件/多件", "dwd_qc_order",
     f"SELECT SUM(is_delivered * qc_missing) / SUM(is_delivered) FROM dwd_qc_order WHERE {W}",
     "%", "越低越好", "预警 > 2.0%", "日更", "已上线",
     "多标签：一单可同时命中多类，五个拆解指标之和 > 品质客诉率，不能相加后当总数用"),
    ("QC102", "质量缺陷率", "L1 结果-品质拆解", "结果", "品质", "QC001",
     "签收后用户反馈商品破损、功能故障、做工差、临期过期的订单占比。",
     "质量缺陷客诉签收单 ÷ 签收订单数", "同 QC001", "同 QC001", "dwd_qc_order",
     f"SELECT SUM(is_delivered * qc_defect) / SUM(is_delivered) FROM dwd_qc_order WHERE {W}",
     "%", "越低越好", "预警 > 1.8%", "日更", "已上线", "外包装压扁但商品完好 → 归包装破损，不算质量缺陷"),
    ("QC103", "货不对板率", "L1 结果-品质拆解", "结果", "品质", "QC001",
     "签收后用户反馈错发、颜色/尺码/型号与描述不符、实物与图片不符的订单占比。",
     "货不对板客诉签收单 ÷ 签收订单数", "同 QC001", "同 QC001", "dwd_qc_order",
     f"SELECT SUM(is_delivered * qc_mismatch) / SUM(is_delivered) FROM dwd_qc_order WHERE {W}",
     "%", "越低越好", "预警 > 1.5%", "日更", "已上线", "\"endereço errado（地址错）\"属于履约问题，规则已排除"),
    ("QC104", "假货投诉率", "L1 结果-品质拆解", "结果", "品质", "QC001",
     "签收后用户反馈非正品、高仿、非原装的订单占比（唯品会\"正品保障\"的核心红线）。",
     "假货客诉签收单 ÷ 签收订单数 × 10000", "同 QC001", "同 QC001", "dwd_qc_order",
     f"SELECT SUM(is_delivered * qc_fake) / SUM(is_delivered) * 10000 FROM dwd_qc_order WHERE {W}",
     "每万单", "越低越好", "红线：单品类 > 30/万单 即启动专项", "日更", "已上线",
     "量小波动大，看趋势用滚动 3 个月；单商家出现 1 单即需人工核实"),
    ("QC105", "包装破损率", "L1 结果-品质拆解", "结果", "品质", "QC001",
     "签收后用户反馈外包装破损、被拆、简陋无防护的订单占比。",
     "包装问题客诉签收单 ÷ 签收订单数", "同 QC001", "同 QC001", "dwd_qc_order",
     f"SELECT SUM(is_delivered * qc_package) / SUM(is_delivered) FROM dwd_qc_order WHERE {W}",
     "%", "越低越好", "预警 > 0.2%", "日更", "已上线", "—"),
    # ------------------------------------------------------------------ L1 体验
    ("EX201", "差评率", "L1 结果-体验", "结果", "体验", "—",
     "已评价订单中 1-2 星订单占比。综合反映品质 + 履约 + 服务，是品质客诉率的\"护栏指标\"。",
     "1-2 星订单数 ÷ 有评价订单数", "全部订单状态（含未签收）；一单一评（取最后一次提交）", "同 QC001", "dwd_qc_order",
     f"SELECT SUM(is_bad) / SUM(has_review) FROM dwd_qc_order WHERE {W}",
     "%", "越低越好", "预警 > 15%", "日更", "已上线",
     "差评率受物流波动影响极大（2018-03 达 22.8%），不能单独用来评价品控成效"),
    ("EX202", "平均评分", "L1 结果-体验", "结果", "体验", "—", "有评价订单的平均星级。",
     "评分之和 ÷ 有评价订单数", "同 EX201", "同 QC001", "dwd_qc_order",
     f"SELECT AVG(review_score) FROM dwd_qc_order WHERE {W}",
     "分", "越高越好", "预警 < 4.0", "日更", "已上线", "均值被极端值拉动，汇报时与差评率一起看"),
    ("EX203", "5 星好评率", "L1 结果-体验", "结果", "体验", "—", "有评价订单中 5 星占比。",
     "5 星订单数 ÷ 有评价订单数", "同 EX201", "同 QC001", "dwd_qc_order",
     f"SELECT SUM(is_five) / SUM(has_review) FROM dwd_qc_order WHERE {W}",
     "%", "越高越好", "—", "日更", "已上线", "—"),
    ("EX204", "差评品质原因占比", "L1 结果-体验", "结果", "体验", "EX201",
     "1-2 星差评中，命中品质类问题的占比。回答\"差评里有多少是品控能管的\"。",
     "品质客诉差评单 ÷ 差评单", "1-2 星", "月、类目、商家", "dwd_qc_order",
     f"SELECT SUM(is_bad * is_quality_complaint) / SUM(is_bad) FROM dwd_qc_order WHERE {W}",
     "%", "看结构", "—", "周更", "已上线",
     "多标签口径（命中即算）；PPT 里的差评原因结构图用的是\"主标签\"互斥口径，两者数值略有差异"),
    # ------------------------------------------------------------------ L2 过程：商品准入
    ("PD301", "动销商品信息完整率", "L2 过程-商品准入", "过程", "商品", "QC103/QC104",
     "有销量的商品中，品类、图片、描述三项信息都完整的商品占比。信息缺失商品的假货投诉率是其他商品的 9 倍。",
     "信息完整的动销商品数 ÷ 动销商品数", "窗口内有订单的商品", "品类、商家", "ods_products + dwd_order_item",
     "SELECT AVG(p.product_category_name IS NOT NULL AND p.product_photos_qty IS NOT NULL "
     "AND p.product_description_lenght IS NOT NULL) FROM ods_products p WHERE p.product_id IN "
     "(SELECT DISTINCT i.product_id FROM dwd_order_item i JOIN dwd_order o ON i.order_id = o.order_id "
     "WHERE o.in_scope = 1 AND o.purchase_month BETWEEN '2018-01' AND '2018-08')",
     "%", "越高越好", "目标 100%（准入卡口）", "周更", "已上线", "按商品去重计算，不是按订单"),
    ("PD302", "信息缺失商品订单占比", "L2 过程-商品准入", "过程", "商品", "PD301",
     "订单主商品缺失品类/图片/描述信息的订单占比。", "主商品信息缺失订单数 ÷ 订单数", "订单主商品 = 订单内售价最高商品",
     "月、商家", "dwd_qc_order",
     f"SELECT AVG(main_category_cn = '未知品类') FROM dwd_qc_order WHERE {W} AND main_product_id IS NOT NULL",
     "%", "越低越好", "预警 > 1%", "周更", "已上线", "—"),
    ("PD303", "高风险品类订单占比", "L2 过程-商品准入", "过程", "商品", "QC001",
     "品质客诉率 ≥ 平台 1.2 倍的品类所产生的订单占比，衡量平台品类结构风险。",
     "高风险品类订单数 ÷ 订单数", "高风险品类名单每季度按全周期数据刷新", "月", "dwd_qc_order + ads_category_quality",
     "SELECT AVG(main_category_cn IN (SELECT category_cn FROM ads_category_quality WHERE qc_rate >= 1.2 * "
     "(SELECT SUM(qc_cnt) / SUM(delivered_cnt) FROM ads_category_quality))) "
     f"FROM dwd_qc_order WHERE {W} AND main_product_id IS NOT NULL",
     "%", "看结构", "—", "月更", "已上线", "名单阈值与平台均值同口径（品类表内的加权平均）"),
    # ------------------------------------------------------------------ L2 过程：商家管理
    ("SP401", "商家发货超时率", "L2 过程-商家管理", "过程", "商家", "QC001/FF501",
     "商家交付物流的时间晚于平台发货 SLA（shipping_limit_date）的订单占比。", "发货超时订单数 ÷ 已出库订单数",
     "已出库订单；商家级考核按\"该商家自己的 SLA\"判定", "月、商家、品类", "dwd_qc_order / dws_seller_month",
     f"SELECT SUM(is_ship_overdue) / SUM(is_shipped) FROM dwd_qc_order WHERE {W}",
     "%", "越低越好", "预警 > 10%", "日更", "已上线", "平台级用订单最晚 SLA；商家级用商家自己的 SLA，两者分母不同"),
    ("SP402", "平均出库时长", "L2 过程-商家管理", "过程", "商家", "SP401", "下单到商家交付物流的平均天数。",
     "Σ(出库时间 − 下单时间) ÷ 已出库订单数", "剔除时间倒挂订单（dq_time_anomaly = 1）", "月、商家", "dwd_qc_order",
     f"SELECT AVG(ship_days) FROM dwd_qc_order WHERE {W} AND is_shipped = 1 AND dq_time_anomaly = 0",
     "天", "越低越好", "预警 > 3 天", "日更", "已上线", "—"),
    ("SP403", "高风险商家数", "L2 过程-商家管理", "过程", "商家", "QC001",
     "近 6 个月签收 ≥ 30 单、平滑品质客诉率 ≥ 平台 2 倍且客诉 ≥ 3 单的商家数。" + SC,
     "COUNT(风险等级 = 高风险)", "签收 ≥ 30 单的参评商家", "月", "ads_seller_scorecard",
     "SELECT COUNT(*) FROM ads_seller_scorecard WHERE risk_level = '高风险'",
     "家", "越低越好", "每月清零目标：进入名单 30 天内整改", "月更", "已上线",
     "用贝叶斯平滑（m=50）后的客诉率判定，避免小样本商家\"1 单客诉=高风险\""),
    ("SP404", "高风险商家客诉贡献度", "L2 过程-商家管理", "过程", "商家", "SP403",
     "高风险商家的品质客诉单占参评商家全部品质客诉单的比例。" + SC, "高风险商家客诉单 ÷ 参评商家客诉单",
     "同 SP403", "—", "ads_seller_scorecard",
     "SELECT SUM(CASE WHEN risk_level = '高风险' THEN qc_cnt ELSE 0 END) / SUM(qc_cnt) FROM ads_seller_scorecard",
     "%", "越低越好", "—", "月更", "已上线", "与\"订单占比\"对照看：订单占比 3.6% vs 客诉占比 12.5%"),
    ("SP405", "商家品质分", "L2 过程-商家管理", "过程", "商家", "SP403",
     "商家级综合得分 0-100：品质客诉率 50% + 差评率 20% + 发货超时率 15% + 延迟签收率 15%，各分项按参评商家分位数打分。" + SC,
     "Σ 权重 × 100 × (1 − PERCENT_RANK)", "同 SP403", "商家", "ads_seller_scorecard",
     "SELECT AVG(quality_score) FROM ads_seller_scorecard",
     "分", "越高越好", "< 40 分进入观察名单", "月更", "已上线",
     "分位制得分，平台均值天然≈50，只用于商家间排序，不适合做平台趋势指标"),
    ("SP406", "参评商家订单覆盖率", "L2 过程-商家管理", "过程", "商家", "SP405",
     "参与品质分评估的商家覆盖的签收订单占比，衡量评分体系的覆盖面。" + SC, "参评商家签收单 ÷ 全部商家签收单",
     "同 SP403", "—", "ads_seller_scorecard + dws_seller_month",
     "SELECT (SELECT SUM(delivered_cnt) FROM ads_seller_scorecard) / SUM(delivered_cnt) FROM dws_seller_month "
     "WHERE purchase_month BETWEEN '2018-03' AND '2018-08'",
     "%", "越高越好", "目标 ≥ 70%", "月更", "已上线", "—"),
    # ------------------------------------------------------------------ L2 过程：仓配履约
    ("FF501", "准时签收率", "L2 过程-仓配履约", "过程", "履约", "EX201",
     "签收日期不晚于承诺送达日期的签收订单占比。", "1 − 延迟签收单 ÷ 签收订单数", "签收订单；按自然日比较",
     "月、省份、商家", "dwd_qc_order",
     f"SELECT 1 - SUM(is_delivered * is_late) / SUM(is_delivered) FROM dwd_qc_order WHERE {W}",
     "%", "越高越好", "预警 < 90%", "日更", "已上线", "承诺日期只有日期部分，必须按日期比较，否则当天送达会被误判延迟"),
    ("FF502", "平均送达时长", "L2 过程-仓配履约", "过程", "履约", "FF501", "下单到用户签收的平均天数。",
     "Σ(签收时间 − 下单时间) ÷ 签收订单数", "剔除时间倒挂订单", "月、省份", "dwd_qc_order",
     f"SELECT AVG(delivery_days) FROM dwd_qc_order WHERE {W} AND is_delivered = 1 AND dq_time_anomaly = 0",
     "天", "越低越好", "—", "日更", "已上线", "最近月份存在右删失（尚未签收的订单未计入），会偏低"),
    ("FF503", "多件订单占比", "L2 过程-仓配履约", "过程", "履约", "QC101", "签收订单中商品行数 ≥ 2 的订单占比。",
     "多件签收单 ÷ 签收订单数", "签收订单", "月、商家、品类", "dwd_qc_order",
     f"SELECT AVG(item_cnt > 1) FROM dwd_qc_order WHERE {W} AND is_delivered = 1",
     "%", "看结构", "—", "周更", "已上线", "—"),
    ("FF504", "多件订单少件率", "L2 过程-仓配履约", "过程", "履约", "QC101",
     "多件签收订单中出现少件/漏发客诉的比例。多件订单占 10% 却贡献 74% 的少件客诉。",
     "多件订单少件客诉单 ÷ 多件签收单", "签收订单且 item_cnt ≥ 2", "月、商家、是否多商家", "dwd_qc_order",
     f"SELECT SUM(qc_missing) / COUNT(*) FROM dwd_qc_order WHERE {W} AND is_delivered = 1 AND item_cnt > 1",
     "%", "越低越好", "目标 ≤ 7%（减半）", "周更", "已上线", "—"),
    ("FF505", "订单取消率", "L2 过程-仓配履约", "过程", "履约", "—", "订单中被取消或因缺货不可履约的占比。",
     "(取消 + 不可用) 订单数 ÷ 订单数", "全部订单", "月、品类", "dwd_qc_order",
     f"SELECT SUM(is_canceled + is_unavailable) / COUNT(*) FROM dwd_qc_order WHERE {W}",
     "%", "越低越好", "预警 > 1.5%", "日更", "已上线", "—"),
    ("FF506", "未收货客诉率", "L2 过程-仓配履约", "过程", "履约", "EX201",
     "用户给出 1-3 星并反馈\"没收到货\"的订单占比。", "未收货客诉单 ÷ 订单数", "全部订单", "月、省份", "dwd_qc_order",
     f"SELECT SUM(fc_not_received) / COUNT(*) FROM dwd_qc_order WHERE {W}",
     "%", "越低越好", "预警 > 5%", "日更", "已上线", "—"),
    # ------------------------------------------------------------------ 待接入
    ("IN601", "入仓质检合格率", "L2 过程-待接入", "过程", "质检", "QC102",
     "供应商商品入仓时，质检合格的批次（或件数）占比。", "质检合格批次 ÷ 质检批次", "入仓质检记录", "供应商、品类、质检项",
     "内部质检系统", None, "%", "越高越好", "—", "日更", "公开数据缺失，需接入",
     "区分\"批次合格率\"与\"件数合格率\"，汇报时写清楚"),
    ("IN602", "抽检覆盖率", "L2 过程-待接入", "过程", "质检", "IN601", "入仓批次中被抽检的比例（高风险供应商应 100%）。",
     "抽检批次 ÷ 入仓批次", "—", "供应商、品类", "内部质检系统", None, "%", "越高越好", "—", "日更", "公开数据缺失，需接入", "—"),
    ("IN603", "问题商品拦截率", "L2 过程-待接入", "过程", "质检", "QC001",
     "质检环节拦截下来的问题件占全部问题件（拦截 + 流出后被客诉）的比例，衡量质检有效性。",
     "质检拦截问题件 ÷ (质检拦截问题件 + 签收后品质客诉件)", "—", "供应商、问题类型", "内部质检系统 + dwd_qc_order",
     None, "%", "越高越好", "—", "周更", "公开数据缺失，需接入", "分子分母问题类型要对齐"),
    ("IN604", "品质退货率", "L1 结果-品质拆解", "结果", "品质", "QC001",
     "因品质原因（质量问题/货不对板/少件/假货）发生的退货件数占签收件数比例；与品质客诉率互为印证。",
     "品质原因退货件 ÷ 签收件", "退货原因 = 品质类", "同 QC001", "售后退货表", None, "%", "越低越好", "—", "日更",
     "公开数据缺失，需接入", "退货原因由用户勾选，存在\"不想要了选质量问题\"的虚报，需结合质检复核结果修正"),
    ("IN605", "客诉首次响应时长", "L2 过程-待接入", "过程", "服务", "EX201", "客诉工单创建到首次响应的时长中位数。",
     "MEDIAN(首次响应时间 − 工单创建时间)", "品质类工单", "客服组、问题类型", "客服工单系统", None, "小时", "越低越好",
     "—", "日更", "公开数据缺失，需接入", "用中位数/P90，不用均值"),
    ("IN606", "高风险商家整改闭环率", "L2 过程-待接入", "过程", "商家", "SP403",
     "进入高风险名单的商家中，30 天内完成整改且复评不再高风险的比例。", "整改闭环商家数 ÷ 应整改商家数", "—", "月",
     "商家治理记录 + ads_seller_scorecard", None, "%", "越高越好", "目标 ≥ 80%", "月更", "公开数据缺失，需接入", "—"),
]

COLS = ["指标编码", "指标名称", "指标层级", "指标类型", "业务域", "父指标", "业务定义", "计算公式", "统计范围 / 过滤条件",
        "可下钻维度", "来源表", "口径 SQL（MySQL 8，可直接执行）", "单位", "方向", "目标 / 预警", "更新频率", "数据可得性",
        "易错点 / 注意事项"]


def compute_baselines():
    vals = []
    for m in M:
        sql = m[11]
        if sql is None:
            vals.append(None)
            continue
        v = query(sql).iloc[0, 0]
        vals.append(float(v) if v is not None else None)
    return vals


def fmt_value(v, unit):
    if v is None:
        return "—"
    if unit == "%":
        return f"{v:.2%}"
    if unit in ("天", "分"):
        return f"{v:.2f}"
    if unit == "每万单":
        return f"{v:.1f}"
    return f"{v:,.0f}"


DIMENSIONS = [
    ("时间", "purchase_month / purchase_date", "下单月 / 下单日", "所有指标统一按下单时间归属（cohort 口径）", "dwd_qc_order"),
    ("一级类目", "main_category_l1", "3C数码、家居家纺等 14 个", "订单主商品（售价最高）所属一级类目", "dim_category"),
    ("品类", "main_category_cn", "73 个品类中文名", "葡语品类映射中文，缺失为\"未知品类\"", "dim_category"),
    ("商家", "seller_id", "3,095 家", "一单多商家时评价归属到每个商家", "dwd_order_seller"),
    ("用户省份", "customer_state", "巴西 27 个州", "用户收货所在州", "ods_customers"),
    ("订单结构", "item_cnt / sku_cnt / seller_cnt", "单件 / 同SKU多件 / 单商家多SKU / 多商家", "由商品行聚合得到", "dwd_order"),
    ("问题类型", "complaint_type", "假货、质量缺陷、货不对板、少件漏发、包装破损、未收到货、物流延迟、服务售后、无文本差评、其他差评",
     "评价文本打标；1-3 星才计客诉", "dwd_review_tag + dim_complaint_rule"),
    ("风险等级", "risk_level", "正常 / 需关注 / 高风险", "商家品质分模型输出", "ads_seller_scorecard"),
]

CHANGELOG = [
    ("v0.1", "初版：评价文本命中品质关键词即记为品质客诉（不限星级）",
     "抽查发现 4-5 星评价里的\"sem defeito（没有瑕疵）\"\"não veio quebrado（没坏）\"被误判"),
    ("v0.2", "品质客诉限定 1-3 星；打标前剔除否定式表述（sem/nenhum/não veio + 缺陷词）",
     "4-5 星误判消除；质量缺陷标签 1-3 星命中数由 1,603 → 1,546"),
    ("v0.3", "\"外包装压扁/破损\"从质量缺陷中剥离，单独归为包装破损",
     "\"箱子破了但商品完好\"不再算质量缺陷，避免高估商品质量问题"),
    ("v0.4", "抽样复核 140 条（7 类 × 20 条，由 AI 逐条阅读葡语原文判定），综合准确率 88.6%；据误判样本修正 8 条规则",
     "未收货类准确率 75%（部分到货被误归未收货）→ 增加\"部分到货\"规则归入少件漏发"),
    ("v1.0", "口径冻结：统一按下单月归属、签收口径、一单一评（最后一次提交）、主标签按优先级互斥",
     "看板、周报、专题分析同源于 dwd_qc_order，数据质量校验 17 项全部通过"),
]


def write_excel(vals):
    wb = Workbook()
    font = "Microsoft YaHei"
    hfill = PatternFill("solid", fgColor="1F3864")
    hfont = Font(name=font, bold=True, color="FFFFFF", size=10)
    bfont = Font(name=font, size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    level_fill = {"L0": "FCE4D6", "L1": "DDEBF7", "L2": "F2F2F2"}

    def header(ws, cols, widths):
        ws.append(cols)
        for i, w in enumerate(widths, 1):
            c = ws.cell(row=1, column=i)
            c.fill, c.font, c.border = hfill, hfont, border
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "C2" if len(cols) > 6 else "A2"
        ws.row_dimensions[1].height = 30

    def body(ws):
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font, c.alignment, c.border = bfont, wrap, border

    # ---- 说明
    ws = wb.active
    ws.title = "说明"
    notes = [
        ("唯品会品控数据分析项目 · 品控指标字典", None),
        ("版本", "v1.0（口径冻结）"),
        ("生成日期", str(date.today())),
        ("数据来源", "Olist Brazilian E-Commerce Public Dataset（巴西真实电商平台脱敏数据，2016-09 ~ 2018-10，约 10 万订单）"),
        ("分析窗口", "2017-01 ~ 2018-08（首尾月份订单过少，已剔除）；本字典\"基准值\"= 2018 年 1-8 月"),
        ("基准值来源", "由 scripts/09_metric_dictionary.py 逐条执行\"口径 SQL\"列得到，非手填"),
        ("口径约定 1", "时间归属：一律按下单月（cohort），保证分子分母为同一批订单"),
        ("口径约定 2", "签收口径：order_status = delivered 且签收时间非空"),
        ("口径约定 3", "一单一评：同一订单多条评价时保留用户最后一次提交的评价"),
        ("口径约定 4", "品质客诉：1-3 星 且 评价文本命中品质类问题（假货/质量缺陷/货不对板/少件漏发/包装破损）"),
        ("口径约定 5", "问题类型为多标签；需要互斥结构（如占比饼图）时使用按优先级取的\"主标签\""),
        ("口径约定 6", "比率类指标只在最后一步计算（DWS 只存计数），禁止\"比率再平均\""),
        ("业务映射", "Olist 商家 ↔ 唯品会供应商/品牌方；用户评价 ↔ 客诉/售后工单；发货 SLA ↔ 供应商发货时效"),
        ("指标层级", "L0 北极星 → L1 结果指标（品质拆解 / 体验）→ L2 过程指标（商品准入 / 商家管理 / 仓配履约 / 待接入）"),
        ("\"待接入\"指标", "入仓质检、退货、客服工单等唯品会内部数据在公开数据中不存在；已定义口径，接入后即可上线"),
    ]
    for k, v in notes:
        ws.append([k, v])
    ws["A1"].font = Font(name=font, bold=True, size=14, color="1F3864")
    for r in range(2, len(notes) + 1):
        ws.cell(row=r, column=1).font = Font(name=font, bold=True, size=10)
        ws.cell(row=r, column=2).font = bfont
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 110

    # ---- 指标体系（树）
    ws = wb.create_sheet("指标体系")
    header(ws, ["指标层级", "业务域", "指标编码", "指标名称", "父指标", "一句话定义", "基准值(2018M1-8)", "方向", "数据可得性"],
           [18, 10, 10, 22, 14, 70, 16, 10, 18])
    for m, v in zip(M, vals):
        ws.append([m[2], m[4], m[0], m[1], m[5], m[6], fmt_value(v, m[12]) if v is not None else "—", m[13], m[16]])
    body(ws)
    for row in ws.iter_rows(min_row=2):
        lv = row[0].value[:2]
        for c in row:
            c.fill = PatternFill("solid", fgColor=level_fill.get(lv, "FFFFFF"))
        if row[8].value != "已上线":
            for c in row:
                c.font = Font(name=font, size=10, color="7F7F7F", italic=True)
    ws.auto_filter.ref = ws.dimensions

    # ---- 指标字典（明细）
    ws = wb.create_sheet("指标字典")
    cols = COLS[:12] + ["基准值(2018M1-8)"] + COLS[12:]
    header(ws, cols, [9, 16, 16, 8, 8, 10, 42, 30, 34, 26, 22, 60, 14, 7, 9, 22, 11, 16, 46])
    for m, v in zip(M, vals):
        ws.append(list(m[:12]) + [fmt_value(v, m[12])] + list(m[12:]))
    body(ws)
    for row in ws.iter_rows(min_row=2):
        row[11].font = Font(name="Consolas", size=9)
        lv = row[2].value[:2]
        for c in row[:2]:
            c.fill = PatternFill("solid", fgColor=level_fill.get(lv, "FFFFFF"))
    ws.auto_filter.ref = ws.dimensions

    # ---- 维度字典
    ws = wb.create_sheet("维度字典")
    header(ws, ["维度", "字段", "取值 / 规模", "口径说明", "来源表"], [12, 30, 50, 46, 30])
    for d in DIMENSIONS:
        ws.append(list(d))
    body(ws)

    # ---- 问题类型字典
    rules = query("SELECT tag_code, tag_name, tag_group, priority FROM dim_complaint_rule ORDER BY priority")
    qa = pd.read_csv(OUT_DIR / "qa" / "tag_validation_sample.csv")
    acc = qa.groupby("primary_tag")["判定（AI 逐条阅读原文）"].mean()
    ws = wb.create_sheet("问题类型字典")
    header(ws, ["类型编码", "问题类型", "分组", "主标签优先级", "是否计入品质客诉", "抽样复核准确率(v0.4，每类20条，AI逐条阅读判定)", "判定示例（葡语原文 → 含义）"],
           [14, 18, 10, 12, 14, 22, 70])
    examples = {
        "fake": "\"recebi um produto falsificado\" → 收到假货", "defect": "\"veio com defeito, não liga\" → 有缺陷、开不了机",
        "mismatch": "\"recebi a cor errada\" → 颜色发错", "missing": "\"comprei 2 e recebi apenas 1\" → 买 2 件只到 1 件",
        "package": "\"caixa veio amassada e violada\" → 箱子压扁被拆", "not_received": "\"ainda não recebi o produto\" → 还没收到货",
        "delay": "\"entrega atrasada\" → 送货延迟", "service": "\"não consigo contato com a loja\" → 联系不上商家",
    }
    for _, r in rules.iterrows():
        a = acc.get(r.tag_code)
        ws.append([r.tag_code, r.tag_name, r.tag_group, int(r.priority), "是" if r.tag_group == "品质类" else "否",
                   f"{a:.0%}" if a is not None and not pd.isna(a) else "未抽样", examples.get(r.tag_code, "")])
    body(ws)

    # ---- 口径变更记录
    ws = wb.create_sheet("口径变更记录")
    header(ws, ["版本", "变更内容", "变更原因 / 影响"], [8, 70, 70])
    for c in CHANGELOG:
        ws.append(list(c))
    body(ws)

    out = ROOT / "docs" / "品控指标字典.xlsx"
    wb.save(out)
    print(f"  → {out.relative_to(ROOT)}")


def write_markdown(vals):
    lines = []
    cur_level = None
    for m, v in zip(M, vals):
        if m[2] != cur_level:
            cur_level = m[2]
            lines += [f"\n### {cur_level}\n", "| 编码 | 指标 | 定义 | 公式 | 基准值 | 目标/预警 |", "|---|---|---|---|---|---|"]
        lines.append(f"| {m[0]} | **{m[1]}** | {m[6]} | {m[7]} | {fmt_value(v, m[12])} | {m[14]} |")
    return "\n".join(lines)


if __name__ == "__main__":
    vals = compute_baselines()
    for m, v in zip(M, vals):
        print(f"  {m[0]} {m[1]:<14} {fmt_value(v, m[12])}")
    write_excel(vals)
    tpl = (ROOT / "docs" / "_templates" / "02_指标体系.md.tpl").read_text(encoding="utf-8")
    doc = tpl.replace("{{METRIC_TABLE}}", write_markdown(vals))
    for m, v in zip(M, vals):
        doc = doc.replace("{" + m[0] + "}", fmt_value(v, m[12]))
    (ROOT / "docs" / "02_指标体系.md").write_text(doc, encoding="utf-8")
    print("  → docs/02_指标体系.md")
