"""Step 15：商家品质月报 Excel 模板（公式驱动 + 数据透视表）。

对应 JD：熟练掌握 Excel（高级公式、透视表）。
  使用说明  怎么用、颜色约定、换数据的步骤
  参数      蓝字为输入：统计起止月、平滑强度、风险阈值；下方为平台基准（SUMIFS）
  月报摘要  核心 KPI + 风险最高的 10 家商家（INDEX/MATCH）
  商家看板  每个商家的签收、客诉、平滑客诉率、差评率、发货超时率、风险等级、排名（SUMIFS / COUNTIFS / IF）
  品类月度  一级类目 × 月份 品质客诉率热力图（SUMIFS + 色阶）
  透视表    LibreOffice 生成的真实数据透视表（行：一级类目；筛选：月份）
  数据      商家 × 月 明细（来自 dws_seller_month，换数据时整表粘贴覆盖）
生成后用 LibreOffice 重算，并与数仓 ads_seller_scorecard 对账（高风险 / 需关注 / 正常 商家数必须一致）。
"""
import subprocess
import time
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common import ROOT, query

OUT = ROOT / "docs" / "商家品质月报模板.xlsx"
FONT = "Microsoft YaHei"
BLUE_INPUT = Font(name=FONT, color="0000FF", bold=True, size=11)
HEAD = Font(name=FONT, bold=True, color="FFFFFF", size=10)
HFILL = PatternFill("solid", fgColor="1F3864")
BODY = Font(name=FONT, size=10)
TITLE = Font(name=FONT, bold=True, size=14, color="1F3864")
NOTE = Font(name=FONT, size=9, color="7F7F7F", italic=True)
YELLOW = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# ------------------------------------------------------------------ 数据
data = query("""
    WITH seller_cat AS (   -- 商家主营一级类目：商品行最多的一级类目
        SELECT seller_id, category_l1 FROM (
            SELECT seller_id, category_l1, ROW_NUMBER() OVER (PARTITION BY seller_id ORDER BY COUNT(*) DESC, category_l1) rn
            FROM dwd_order_item GROUP BY seller_id, category_l1) t WHERE rn = 1)
    SELECT m.purchase_month, CAST(REPLACE(m.purchase_month, '-', '') AS UNSIGNED) AS month_key,
           m.seller_id, LEFT(m.seller_id, 8) AS seller_short, COALESCE(c.category_l1, '未知') AS category_l1,
           m.delivered_cnt, m.qc_cnt, m.review_cnt, m.bad_cnt, m.shipped_cnt, m.ship_overdue_cnt, m.late_cnt
    FROM dws_seller_month m LEFT JOIN seller_cat c ON c.seller_id = m.seller_id
    ORDER BY m.purchase_month, m.seller_id""")
for c in ["month_key", "delivered_cnt", "qc_cnt", "review_cnt", "bad_cnt", "shipped_cnt", "ship_overdue_cnt", "late_cnt"]:
    data[c] = data[c].astype(int)
N = len(data) + 1                                   # 数据最后一行行号
sellers = (data.groupby(["seller_id", "seller_short", "category_l1"]).delivered_cnt.sum().reset_index()
           .query("delivered_cnt >= 10").sort_values("delivered_cnt", ascending=False))
S0, S1 = 6, 6 + len(sellers) - 1                    # 商家看板数据行范围
months = sorted(data.purchase_month.unique())
cats = sorted(data.category_l1.unique(), key=lambda x: (x == "未知", x))

wb = Workbook()


def header(ws, row, cols, widths=None):
    for i, t in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.border = HEAD, HFILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


D = "'数据'"
rng = lambda col: f"{D}!${col}$2:${col}${N}"
in_period = f"{rng('B')},\">=\"&'参数'!$B$4,{rng('B')},\"<=\"&'参数'!$B$5"

# ------------------------------------------------------------------ 使用说明
ws = wb.active
ws.title = "使用说明"
lines = [
    ("商家品质月报模板", TITLE),
    ("用途：每月（或每周）更新一次，给品控 / 商家管理团队看平台品质客诉率、风险商家名单和品类热力图。", BODY),
    ("", BODY),
    ("怎么用", Font(name=FONT, bold=True, size=11)),
    ("1. 在「参数」页修改蓝色字体的输入格：统计起止月份（格式 201803）、平滑强度、风险阈值。其他页全部自动重算。", BODY),
    ("2. 「月报摘要」直接复制进周报 / 月报；「商家看板」可按风险等级筛选、按排名排序。", BODY),
    ("3. 换数据：把新的商家 × 月明细（字段顺序与「数据」页一致）粘贴覆盖「数据」页，再在「透视表」页右键 → 刷新。", BODY),
    ("4. 本模板底层是月度数据；改成周报时，把「数据」页的月份键换成周键（如 201814 表示 2018 年第 14 周）即可，公式不用改。", BODY),
    ("", BODY),
    ("颜色约定", Font(name=FONT, bold=True, size=11)),
    ("蓝色字体 + 浅黄底 = 可以修改的输入 / 假设；黑色字体 = 公式，不要手改。", BODY),
    ("风险等级：红 = 高风险，橙 = 需关注，灰 = 未参评（签收单不足）。", BODY),
    ("", BODY),
    ("口径（与指标字典 v1.0 一致）", Font(name=FONT, bold=True, size=11)),
    ("品质客诉率 = 品质客诉单 ÷ 签收单；按下单月归属；一单多商家时评价计入每个商家。", BODY),
    ("平滑客诉率 =（客诉单 + m × 平台客诉率）÷（签收单 + m），用于小样本商家排序，m 默认 50。", BODY),
    ("高风险 = 签收单 ≥ 最小签收单 且 客诉单 ≥ 3 且 平滑客诉率 ≥ 平台 × 2；需关注 = 平滑客诉率 ≥ 平台 × 1.5 或 差评率 ≥ 平台 × 1.5。", BODY),
    ("", BODY),
    ("用到的函数：SUMIFS、COUNTIFS、INDEX + MATCH、IF / AND / OR、IFERROR、N；条件格式（色阶、数据条、公式规则）；数据透视表。", BODY),
    ("数据来源：Olist 巴西电商公开数据 → 项目数仓 dws_seller_month（scripts/15_excel_template.py 生成）。", NOTE),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = f
ws.column_dimensions["A"].width = 120

# ------------------------------------------------------------------ 参数
ws = wb.create_sheet("参数")
ws["A1"], ws["A1"].font = "参数与平台基准", TITLE
ws["A2"], ws["A2"].font = "蓝字浅黄底为输入；修改后所有页面自动重算", NOTE
params = [
    ("A4", "统计起始月（月份键）", "B4", 201803),
    ("A5", "统计截止月（月份键）", "B5", 201808),
    ("A6", "平滑强度 m", "B6", 50),
    ("A7", "高风险：平滑客诉率 ≥ 平台 ×", "B7", 2),
    ("A8", "需关注：平滑客诉率 / 差评率 ≥ 平台 ×", "B8", 1.5),
    ("A9", "参评门槛：签收单 ≥", "B9", 30),
    ("A10", "高风险门槛：客诉单 ≥", "B10", 3),
]
for a, label, b, v in params:
    ws[a], ws[a].font = label, BODY
    ws[b] = v
    ws[b].font, ws[b].fill, ws[b].border = BLUE_INPUT, YELLOW, BORDER
ws["A12"], ws["A12"].font = "平台基准（公式）", Font(name=FONT, bold=True, size=11)
base = [
    ("A13", "平台签收单", "B13", f"=SUMIFS({rng('F')},{in_period})", "#,##0"),
    ("A14", "平台品质客诉单", "B14", f"=SUMIFS({rng('G')},{in_period})", "#,##0"),
    ("A15", "平台品质客诉率", "B15", "=IFERROR(B14/B13,0)", "0.00%"),
    ("A16", "平台已评价单", "B16", f"=SUMIFS({rng('H')},{in_period})", "#,##0"),
    ("A17", "平台差评单", "B17", f"=SUMIFS({rng('I')},{in_period})", "#,##0"),
    ("A18", "平台差评率", "B18", "=IFERROR(B17/B16,0)", "0.00%"),
]
for a, label, b, f, fmt in base:
    ws[a], ws[a].font = label, BODY
    ws[b] = f
    ws[b].font, ws[b].number_format, ws[b].border = BODY, fmt, BORDER
ws["A20"], ws["A20"].font = "月份键说明：201803 = 2018 年 3 月；数据范围 201701 ~ 201808", NOTE
ws.column_dimensions["A"].width = 38
ws.column_dimensions["B"].width = 16

# ------------------------------------------------------------------ 商家看板
ws = wb.create_sheet("商家看板")
ws["A1"], ws["A1"].font = "商家品质看板", TITLE
ws["A2"] = '="统计期："&\'参数\'!B4&" ~ "&\'参数\'!B5&"　｜　平台品质客诉率 "&TEXT(\'参数\'!B15,"0.00%")&"　｜　仅列出全周期签收 ≥ 10 单的商家"'
ws["A2"].font = NOTE
ws["A3"] = '="高风险 "&COUNTIF($L$6:$L$' + str(S1) + ',"高风险")&" 家　需关注 "&COUNTIF($L$6:$L$' + str(S1) + ',"需关注")&" 家　正常 "&COUNTIF($L$6:$L$' + str(S1) + ',"正常")&" 家　未参评 "&COUNTIF($L$6:$L$' + str(S1) + ',"未参评")&" 家"'
ws["A3"].font = Font(name=FONT, bold=True, size=10)
cols = ["商家（前 8 位）", "商家 ID", "主营类目", "签收单", "品质客诉单", "品质客诉率", "平滑客诉率", "已评价单", "差评率",
        "已出库单", "发货超时率", "风险等级", "排名"]
header(ws, 5, cols, [14, 36, 11, 9, 10, 10, 10, 9, 9, 9, 10, 9, 7])
P = "'参数'"
for i, r in enumerate(sellers.itertuples(index=False), start=S0):
    ws.cell(row=i, column=1, value=r.seller_short)
    ws.cell(row=i, column=2, value=r.seller_id)
    ws.cell(row=i, column=3, value=r.category_l1)
    by = f"{rng('C')},$B{i},{in_period}"
    ws.cell(row=i, column=4, value=f"=SUMIFS({rng('F')},{by})")
    ws.cell(row=i, column=5, value=f"=SUMIFS({rng('G')},{by})")
    ws.cell(row=i, column=6, value=f'=IF(D{i}>0,E{i}/D{i},"")')
    ws.cell(row=i, column=7, value=f"=(E{i}+{P}!$B$6*{P}!$B$15)/(D{i}+{P}!$B$6)")
    ws.cell(row=i, column=8, value=f"=SUMIFS({rng('H')},{by})")
    ws.cell(row=i, column=9, value=f'=IF(H{i}>0,SUMIFS({rng("I")},{by})/H{i},"")')
    ws.cell(row=i, column=10, value=f"=SUMIFS({rng('J')},{by})")
    ws.cell(row=i, column=11, value=f'=IF(J{i}>0,SUMIFS({rng("K")},{by})/J{i},"")')
    ws.cell(row=i, column=12, value=(f'=IF(D{i}<{P}!$B$9,"未参评",IF(AND(G{i}>={P}!$B$7*{P}!$B$15,E{i}>={P}!$B$10),"高风险",'
                                     f'IF(OR(G{i}>={P}!$B$8*{P}!$B$15,N(I{i})>={P}!$B$8*{P}!$B$18),"需关注","正常")))'))
    # 唯一排名：参评商家按平滑客诉率从高到低；并列时按行号先后
    ws.cell(row=i, column=13, value=(f'=IF(L{i}="未参评","",COUNTIFS($L${S0}:$L${S1},"<>未参评",$G${S0}:$G${S1},">"&G{i})'
                                     f'+COUNTIFS($L${S0}:L{i},"<>未参评",$G${S0}:G{i},G{i}))'))
    for col, fmt in [(4, "#,##0"), (5, "#,##0"), (6, "0.0%"), (7, "0.0%"), (8, "#,##0"), (9, "0.0%"), (10, "#,##0"), (11, "0.0%")]:
        ws.cell(row=i, column=col).number_format = fmt
    for col in range(1, 14):
        ws.cell(row=i, column=col).font = BODY
ws.freeze_panes = "D6"
ws.auto_filter.ref = f"A5:M{S1}"
ws.conditional_formatting.add(f"L{S0}:L{S1}", CellIsRule(operator="equal", formula=['"高风险"'],
                              fill=PatternFill("solid", fgColor="F4CCCC"), font=Font(color="B42323", bold=True)))
ws.conditional_formatting.add(f"L{S0}:L{S1}", CellIsRule(operator="equal", formula=['"需关注"'],
                              fill=PatternFill("solid", fgColor="FCE5CD"), font=Font(color="8A5800", bold=True)))
ws.conditional_formatting.add(f"L{S0}:L{S1}", CellIsRule(operator="equal", formula=['"未参评"'], font=Font(color="999999")))
ws.conditional_formatting.add(f"G{S0}:G{S1}", DataBarRule(start_type="num", start_value=0, end_type="max", color="5B9BD5"))

# ------------------------------------------------------------------ 月报摘要
ws = wb.create_sheet("月报摘要", 1)
K = "'商家看板'"
ws["A1"], ws["A1"].font = "商家品质月报摘要", TITLE
ws["A2"] = '="统计期："&\'参数\'!B4&" ~ "&\'参数\'!B5'
ws["A2"].font = NOTE
kpis = [
    ("签收订单", f"={P}!B13", "#,##0"),
    ("品质客诉单", f"={P}!B14", "#,##0"),
    ("品质客诉率", f"={P}!B15", "0.00%"),
    ("差评率", f"={P}!B18", "0.00%"),
    ("高风险商家数", f'=COUNTIF({K}!$L${S0}:$L${S1},"高风险")', "0"),
    ("需关注商家数", f'=COUNTIF({K}!$L${S0}:$L${S1},"需关注")', "0"),
    ("高风险商家客诉占全平台", f'=IFERROR(SUMIFS({K}!$E${S0}:$E${S1},{K}!$L${S0}:$L${S1},"高风险")/{P}!B14,0)', "0.0%"),
    ("高风险商家签收占全平台", f'=IFERROR(SUMIFS({K}!$D${S0}:$D${S1},{K}!$L${S0}:$L${S1},"高风险")/{P}!B13,0)', "0.0%"),
]
for j, (label, f, fmt) in enumerate(kpis):
    col = 1 + (j % 4) * 2
    row = 4 + (j // 4) * 3
    a = ws.cell(row=row, column=col, value=label)
    a.font = Font(name=FONT, size=9, color="595959")
    b = ws.cell(row=row + 1, column=col, value=f)
    b.font, b.number_format = Font(name=FONT, size=18, bold=True, color="1F3864"), fmt
ws["A11"], ws["A11"].font = "风险最高的 10 家参评商家（按平滑客诉率）", Font(name=FONT, bold=True, size=11)
header(ws, 12, ["排名", "商家", "主营类目", "签收单", "品质客诉单", "品质客诉率", "平滑客诉率", "差评率", "风险等级"],
       [8, 14, 12, 10, 12, 12, 12, 10, 10])
src_cols = {2: "A", 3: "C", 4: "D", 5: "E", 6: "F", 7: "G", 8: "I", 9: "L"}
fmts = {4: "#,##0", 5: "#,##0", 6: "0.0%", 7: "0.0%", 8: "0.0%"}
for k in range(1, 11):
    r = 12 + k
    ws.cell(row=r, column=1, value=k).font = BODY
    for col, sc in src_cols.items():
        c = ws.cell(row=r, column=col, value=f'=IFERROR(INDEX({K}!${sc}${S0}:${sc}${S1},MATCH($A{r},{K}!$M${S0}:$M${S1},0)),"")')
        c.font, c.border = BODY, BORDER
        if col in fmts:
            c.number_format = fmts[col]
ws.conditional_formatting.add("I13:I22", CellIsRule(operator="equal", formula=['"高风险"'],
                              fill=PatternFill("solid", fgColor="F4CCCC"), font=Font(color="B42323", bold=True)))
ws.conditional_formatting.add("I13:I22", CellIsRule(operator="equal", formula=['"需关注"'],
                              fill=PatternFill("solid", fgColor="FCE5CD"), font=Font(color="8A5800", bold=True)))

# ------------------------------------------------------------------ 品类月度
ws = wb.create_sheet("品类月度")
ws["A1"], ws["A1"].font = "一级类目 × 月份 品质客诉率（色阶：越红越高）", TITLE
ws["A2"], ws["A2"].font = "按商家主营类目汇总；最右列为全周期合计", NOTE
header(ws, 4, ["一级类目"] + [m[2:] for m in months] + ["全周期"], [12] + [7] * len(months) + [9])
for i, cat in enumerate(cats, start=5):
    ws.cell(row=i, column=1, value=cat).font = BODY
    for j, m in enumerate(months, start=2):
        f = (f'=IFERROR(SUMIFS({rng("G")},{rng("E")},$A{i},{rng("A")},"{m}")/'
             f'SUMIFS({rng("F")},{rng("E")},$A{i},{rng("A")},"{m}"),"")')
        c = ws.cell(row=i, column=j, value=f)
        c.number_format, c.font = "0.0%", Font(name=FONT, size=9)
    c = ws.cell(row=i, column=len(months) + 2, value=f'=IFERROR(SUMIFS({rng("G")},{rng("E")},$A{i})/SUMIFS({rng("F")},{rng("E")},$A{i}),"")')
    c.number_format, c.font = "0.0%", Font(name=FONT, size=9, bold=True)
last = get_column_letter(len(months) + 1)
ws.conditional_formatting.add(f"B5:{last}{4 + len(cats)}",
                              ColorScaleRule(start_type="num", start_value=0.02, start_color="FFFFFF",
                                             mid_type="num", mid_value=0.05, mid_color="FCE5CD",
                                             end_type="num", end_value=0.10, end_color="E06666"))
ws.freeze_panes = "B5"

# ------------------------------------------------------------------ 透视表（占位，LibreOffice 生成）
ws = wb.create_sheet("透视表")
ws["A1"], ws["A1"].font = "数据透视表：一级类目 × 签收 / 品质客诉 / 差评（顶部筛选器可选月份）", TITLE

# ------------------------------------------------------------------ 数据
ws = wb.create_sheet("数据")
cols = ["月份", "月份键", "商家ID", "商家简称", "主营一级类目", "签收单", "品质客诉单", "已评价单", "差评单", "已出库单", "发货超时单", "延迟签收单"]
header(ws, 1, cols, [9, 9, 36, 11, 12, 8, 10, 9, 8, 9, 10, 10])
for r in data.itertuples(index=False):
    ws.append(list(r))
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:L{N}"

OUT.parent.mkdir(exist_ok=True)
tmp = OUT.with_name("_tmp_template.xlsx")
wb.save(tmp)
print(f"  openpyxl 写入完成：数据 {len(data):,} 行，商家 {len(sellers):,} 家")

# ------------------------------------------------------------------ LibreOffice：插入透视表 + 重算 + 另存为 xlsx
UNO_SCRIPT = r'''
import sys, time, uno
from com.sun.star.beans import PropertyValue
from com.sun.star.table import CellAddress, CellRangeAddress
def pv(n, v):
    p = PropertyValue(); p.Name = n; p.Value = v; return p
src_path, out_path = sys.argv[1], sys.argv[2]
local = uno.getComponentContext()
resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
for _ in range(60):
    try:
        ctx = resolver.resolve("uno:socket,host=localhost,port=2099;urp;StarOffice.ComponentContext"); break
    except Exception:
        time.sleep(1)
desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(src_path), "_blank", 0, (pv("Hidden", True),))
sheets = doc.Sheets
names = list(sheets.ElementNames)
data = sheets.getByName("数据")
cur = data.createCursor(); cur.gotoEndOfUsedArea(False)
end_row, end_col = cur.RangeAddress.EndRow, cur.RangeAddress.EndColumn
src = CellRangeAddress(); src.Sheet = names.index("数据"); src.StartColumn = 0; src.StartRow = 0
src.EndColumn = end_col; src.EndRow = end_row
dst = sheets.getByName("透视表")
tables = dst.getDataPilotTables()
ORI = lambda x: uno.Enum("com.sun.star.sheet.DataPilotFieldOrientation", x)
SUM = uno.Enum("com.sun.star.sheet.GeneralFunction", "SUM")
def make(name, row_field, page_field, data_fields, anchor_row):
    desc = tables.createDataPilotDescriptor()
    desc.setSourceRange(src)
    fields = desc.getDataPilotFields()
    for i in range(fields.getCount()):
        f = fields.getByIndex(i)
        if f.Name == row_field: f.Orientation = ORI("ROW")
        elif page_field and f.Name == page_field: f.Orientation = ORI("PAGE")
        elif f.Name in data_fields:
            f.Orientation = ORI("DATA"); f.Function = SUM
    try:
        desc.getDataLayoutField().Orientation = ORI("COLUMN")   # 多个数据字段横向排列
    except Exception:
        pass
    tables.insertNewByName(name, CellAddress(names.index("透视表"), 0, anchor_row), desc)
    try:   # 数据字段显示名改成中文
        t = tables.getByName(name)
        dfs = t.getDataFields()
        for i in range(dfs.getCount()):
            f = dfs.getByIndex(i)
            f.setPropertyValue("LayoutName", "合计：" + f.getName().split(" - ")[-1].strip())
    except Exception as e:
        print("layoutname skipped", e)
make("类目透视", "主营一级类目", "月份", ["签收单", "品质客诉单", "差评单"], 3)
make("月度透视", "月份", None, ["签收单", "品质客诉单"], 24)
doc.calculateAll()
doc.storeToURL(uno.systemPathToFileUrl(out_path), (pv("FilterName", "Calc MS Excel 2007 XML"),))
doc.close(True)
print("uno ok")
'''
script = Path("/tmp/_uno_pivot.py")
script.write_text(UNO_SCRIPT, encoding="utf-8")
office = subprocess.Popen(["soffice", "--headless", "--invisible", "--norestore", "--nologo",
                           "--accept=socket,host=localhost,port=2099;urp;"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
try:
    res = subprocess.run(["/usr/bin/python3", str(script), str(tmp), str(OUT)], capture_output=True, text=True, timeout=600)
    print("  LibreOffice：", res.stdout.strip() or res.stderr.strip()[-500:])
finally:
    office.terminate()
    time.sleep(2)
    office.kill()
tmp.unlink(missing_ok=True)

# ------------------------------------------------------------------ 对账：模板算出的商家分层 = 数仓评分卡
v = load_workbook(OUT, data_only=True)
ks = v["商家看板"]
tiers = pd.Series([ks.cell(row=i, column=12).value for i in range(S0, S1 + 1)]).value_counts()
sc = query("SELECT risk_level, COUNT(*) n FROM ads_seller_scorecard GROUP BY 1").set_index("risk_level").n
ok = all(int(tiers.get(t, 0)) == int(sc.get(t, 0)) for t in ["高风险", "需关注", "正常"])
print(f"  对账：模板 高风险 {tiers.get('高风险', 0)} / 需关注 {tiers.get('需关注', 0)} / 正常 {tiers.get('正常', 0)}"
      f"；数仓 {sc.get('高风险', 0)} / {sc.get('需关注', 0)} / {sc.get('正常', 0)} → {'一致' if ok else '不一致'}")
print(f"  平台品质客诉率（模板）{v['参数']['B15'].value:.4%}")
if not ok:
    raise SystemExit("Excel 模板与数仓评分卡不一致")
print(f"  → {OUT.relative_to(ROOT)}")
